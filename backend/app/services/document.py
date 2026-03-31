import logging
from datetime import datetime, timezone
from pathlib import Path
from typing import Optional
from uuid import UUID

from jinja2 import Environment, FileSystemLoader
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.models.audit_entry import AuditEntry
from app.models.base import (
    ActorType,
    AuditAction,
    DocumentType,
    FileFormat,
)
from app.models.case import Case
from app.models.client import Client
from app.models.document import Document
from app.models.recommendation import Recommendation
from app.models.user import User

logger = logging.getLogger(__name__)

GENERATED_DOCS_ROOT = Path(__file__).resolve().parent.parent.parent / "generated_documents"


class DocumentService:
    """Generates compliance-ready documents from recommendations."""

    def __init__(self, db: AsyncSession):
        self.db = db
        prompts_dir = Path(__file__).resolve().parent.parent / "prompts"
        self.jinja_env = Environment(
            loader=FileSystemLoader(str(prompts_dir)),
            autoescape=True,
        )

    async def generate_recommendation_pack(
        self,
        recommendation_id: UUID,
        generated_by: UUID,
        organization_id: UUID,
        file_format: FileFormat = FileFormat.DOCX,
    ) -> Document:
        """Generate a recommendation pack document (DOCX or PDF)."""

        # Load recommendation with evidences
        result = await self.db.execute(
            select(Recommendation)
            .options(selectinload(Recommendation.evidences))
            .where(Recommendation.id == recommendation_id)
        )
        recommendation = result.scalar_one_or_none()
        if not recommendation:
            raise ValueError(f"Recommendation {recommendation_id} not found")

        # Load case
        result = await self.db.execute(
            select(Case).where(
                Case.id == recommendation.case_id,
                Case.organization_id == organization_id,
            )
        )
        case = result.scalar_one_or_none()
        if not case:
            raise ValueError("Case not found or does not belong to organization")

        # Load client
        result = await self.db.execute(
            select(Client).where(Client.id == case.client_id)
        )
        client = result.scalar_one_or_none()
        if not client:
            raise ValueError("Client not found")

        # Load advisor (assigned_to)
        result = await self.db.execute(
            select(User).where(User.id == case.assigned_to)
        )
        advisor = result.scalar_one_or_none()

        now = datetime.now(timezone.utc)
        generated_at = now.strftime("%Y-%m-%d %H:%M")

        # Build template context
        template_ctx = self._build_template_context(
            recommendation, case, client, advisor, generated_at
        )

        # Render HTML
        template = self.jinja_env.get_template("recommendation_pack.html")
        html_content = template.render(**template_ctx)

        # Create output directory
        doc_dir = GENERATED_DOCS_ROOT / str(organization_id) / str(case.id)
        doc_dir.mkdir(parents=True, exist_ok=True)

        # Generate filename
        safe_name = client.name.replace(" ", "_")[:30]
        filename = f"rekommendation_{safe_name}_v{recommendation.version}.{file_format.value}"
        file_path = doc_dir / filename

        # Generate file
        if file_format == FileFormat.PDF:
            self._generate_pdf(html_content, file_path)
        else:
            self._generate_docx(html_content, template_ctx, file_path)

        # Get next document version for this recommendation
        from sqlalchemy import func

        version_result = await self.db.execute(
            select(func.coalesce(func.max(Document.version), 0) + 1).where(
                Document.recommendation_id == recommendation_id
            )
        )
        next_version = version_result.scalar()

        # Create Document record
        document = Document(
            case_id=case.id,
            recommendation_id=recommendation_id,
            document_type=DocumentType.RECOMMENDATION_PACK,
            title=f"Rekommendationsunderlag — {client.name} v{recommendation.version}",
            template_id="recommendation_pack.html",
            file_path=str(file_path),
            file_format=file_format,
            generated_by=generated_by,
            version=next_version,
        )
        self.db.add(document)
        await self.db.flush()

        # Audit entry
        audit = AuditEntry(
            case_id=case.id,
            action=AuditAction.DOCUMENT_GENERATED,
            actor_id=generated_by,
            actor_type=ActorType.USER,
            details={
                "document_id": str(document.id),
                "document_type": DocumentType.RECOMMENDATION_PACK.value,
                "template_id": "recommendation_pack.html",
                "recommendation_id": str(recommendation_id),
                "recommendation_version": recommendation.version,
                "file_format": file_format.value,
                "file_path": str(file_path),
            },
        )
        self.db.add(audit)
        await self.db.flush()
        await self.db.refresh(document)

        logger.info(
            "Generated %s document for recommendation %s → %s",
            file_format.value,
            recommendation_id,
            file_path,
        )
        return document

    def _build_template_context(
        self,
        recommendation: Recommendation,
        case: Case,
        client: Client,
        advisor: Optional[User],
        generated_at: str,
    ) -> dict:
        """Build the Jinja2 template context from domain objects."""

        # Convert reasoning_chain items to dicts if needed
        reasoning_chain = recommendation.reasoning_chain or []
        if reasoning_chain and not isinstance(reasoning_chain[0], dict):
            reasoning_chain = [vars(s) for s in reasoning_chain]

        assumptions = recommendation.assumptions or []
        if assumptions and not isinstance(assumptions[0], dict):
            assumptions = [vars(a) for a in assumptions]

        scenarios = recommendation.scenarios or []
        if scenarios and not isinstance(scenarios[0], dict):
            scenarios = [vars(s) for s in scenarios]

        evidences = []
        for ev in (recommendation.evidences or []):
            evidences.append({
                "source_type": ev.source_type.value if hasattr(ev.source_type, 'value') else ev.source_type,
                "source_reference": ev.source_reference,
                "content_snippet": ev.content_snippet,
                "relevance_explanation": ev.relevance_explanation,
                "confidence": float(ev.confidence),
            })

        return {
            "client": {
                "name": client.name,
                "date_of_birth": str(client.date_of_birth),
                "employment_status": client.employment_status.value if hasattr(client.employment_status, 'value') else client.employment_status,
                "employer_name": client.employer_name,
                "collective_agreement": client.collective_agreement.value if hasattr(client.collective_agreement, 'value') else client.collective_agreement,
                "annual_income": float(client.annual_income) if client.annual_income else None,
                "desired_retirement_age": client.desired_retirement_age,
                "risk_profile": client.risk_profile.value if client.risk_profile and hasattr(client.risk_profile, 'value') else client.risk_profile,
            },
            "case": {
                "id": str(case.id),
                "title": case.title,
                "case_type": case.case_type.value if hasattr(case.case_type, 'value') else case.case_type,
                "summary": case.summary,
            },
            "recommendation": {
                "summary": recommendation.summary,
                "recommendation_type": recommendation.recommendation_type.value if hasattr(recommendation.recommendation_type, 'value') else recommendation.recommendation_type,
                "version": recommendation.version,
                "status": recommendation.status.value if hasattr(recommendation.status, 'value') else recommendation.status,
                "suitability_score": float(recommendation.suitability_score) if recommendation.suitability_score else None,
                "reasoning_chain": reasoning_chain,
                "assumptions": assumptions,
                "scenarios": scenarios,
            },
            "evidences": evidences,
            "advisor": {
                "name": advisor.name if advisor else "Okänd",
                "email": advisor.email if advisor else "",
            },
            "generated_at": generated_at,
        }

    def _generate_pdf(self, html_content: str, output_path: Path) -> None:
        """Generate PDF from HTML using WeasyPrint."""
        from weasyprint import HTML

        HTML(string=html_content).write_pdf(str(output_path))

    def _generate_docx(
        self, html_content: str, ctx: dict, output_path: Path
    ) -> None:
        """Generate DOCX from template context using python-docx."""
        from docx import Document as DocxDocument
        from docx.shared import Pt, Inches, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.table import WD_TABLE_ALIGNMENT

        doc = DocxDocument()

        style = doc.styles["Normal"]
        style.font.name = "Calibri"
        style.font.size = Pt(10.5)
        style.paragraph_format.space_after = Pt(4)

        client = ctx["client"]
        case = ctx["case"]
        rec = ctx["recommendation"]
        evidences = ctx["evidences"]
        advisor = ctx["advisor"]
        generated_at = ctx["generated_at"]

        # --- Title ---
        title_para = doc.add_heading("Rekommendationsunderlag", level=0)
        title_para.runs[0].font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)

        meta = doc.add_paragraph()
        meta_run = meta.add_run(f"{client['name']} · {case['title']} · {generated_at}")
        meta_run.font.size = Pt(9.5)
        meta_run.font.color.rgb = RGBColor(0x71, 0x80, 0x96)

        # --- Helper ---
        def add_section(number: int, title: str):
            h = doc.add_heading(f"{number}. {title}", level=2)
            for run in h.runs:
                run.font.color.rgb = RGBColor(0x2C, 0x52, 0x82)

        def add_kv_table(pairs: list[tuple[str, str]]):
            table = doc.add_table(rows=0, cols=2)
            table.alignment = WD_TABLE_ALIGNMENT.LEFT
            table.style = "Light List Accent 1"
            for label, value in pairs:
                if value is None or value == "" or value == "None":
                    continue
                row = table.add_row()
                row.cells[0].text = label
                row.cells[1].text = str(value)
                row.cells[0].paragraphs[0].runs[0].font.bold = True if row.cells[0].paragraphs[0].runs else None

        # --- 1. Kundsammanfattning ---
        add_section(1, "Kundsammanfattning")
        income_str = f"{client['annual_income']:,.0f} SEK" if client.get("annual_income") else None
        add_kv_table([
            ("Namn", client["name"]),
            ("Födelsedatum", client["date_of_birth"]),
            ("Anställningsstatus", client["employment_status"]),
            ("Arbetsgivare", client.get("employer_name")),
            ("Kollektivavtal", client["collective_agreement"]),
            ("Årsinkomst", income_str),
            ("Önskad pensionsålder", f"{client['desired_retirement_age']} år" if client.get("desired_retirement_age") else None),
            ("Riskprofil", client.get("risk_profile")),
        ])

        # --- 2. Behovsanalys ---
        add_section(2, "Behovsanalys")
        doc.add_paragraph(rec["summary"])
        needs_steps = rec["reasoning_chain"][:2] if len(rec["reasoning_chain"]) > 2 else rec["reasoning_chain"]
        for step in needs_steps:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            step_label = step.get("title") or f"Steg {step['step']}"
            bold_run = p.add_run(f"{step_label}: ")
            bold_run.bold = True
            p.add_run(step["description"])
            # Cited texts as indented blockquotes
            for ct in step.get("cited_texts", []):
                cq = doc.add_paragraph()
                cq.paragraph_format.left_indent = Inches(0.5)
                cq_run = cq.add_run(f"\u201c{ct['text']}\u201d")
                cq_run.italic = True
                cq_run.font.size = Pt(9)
                cq_run.font.color.rgb = RGBColor(0x71, 0x80, 0x96)
                if ct.get("source_title"):
                    src_run = cq.add_run(f" \u2014 {ct['source_title']}")
                    src_run.font.size = Pt(8.5)
                    src_run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
            if step.get("conclusion"):
                conclusion_p = doc.add_paragraph()
                conclusion_p.paragraph_format.left_indent = Inches(0.3)
                conclusion_run = conclusion_p.add_run(step["conclusion"])
                conclusion_run.italic = True

        # --- 3. Marknads- och produktanalys ---
        add_section(3, "Marknads- och produktanalys")
        scenarios = rec.get("scenarios") or []
        if scenarios:
            # Collect all outcome keys
            all_keys = []
            for s in scenarios:
                for k in s.get("projected_outcome", {}).keys():
                    if k not in all_keys:
                        all_keys.append(k)

            headers = ["Alternativ", "Beskrivning"] + [k.replace("_", " ").capitalize() for k in all_keys]
            table = doc.add_table(rows=1, cols=len(headers))
            table.style = "Light List Accent 1"
            for i, h_text in enumerate(headers):
                table.rows[0].cells[i].text = h_text

            for s in scenarios:
                row = table.add_row()
                row.cells[0].text = s["name"]
                row.cells[1].text = s["description"]
                for j, k in enumerate(all_keys):
                    row.cells[2 + j].text = str(s.get("projected_outcome", {}).get(k, "—"))
        else:
            doc.add_paragraph("Ingen scenariojämförelse tillgänglig.")

        # --- 4. Rekommendation ---
        add_section(4, "Rekommendation")
        doc.add_paragraph(rec["summary"])
        add_kv_table([
            ("Typ", rec["recommendation_type"]),
            ("Version", str(rec["version"])),
            ("Status", rec["status"]),
        ])

        # --- 5. Motivering ---
        add_section(5, "Motivering")
        doc.add_heading("Resonemangskedja", level=3)
        for step in rec["reasoning_chain"]:
            desc = step.get("description", "")
            if "Kostnadsinformation" in desc or "Intressekonflikt" in desc:
                continue
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            step_label = step.get("title") or f"Steg {step['step']}"
            bold_run = p.add_run(f"{step_label}: ")
            bold_run.bold = True
            p.add_run(step["description"])
            # Cited texts
            for ct in step.get("cited_texts", []):
                cq = doc.add_paragraph()
                cq.paragraph_format.left_indent = Inches(0.5)
                cq_run = cq.add_run(f"\u201c{ct['text']}\u201d")
                cq_run.italic = True
                cq_run.font.size = Pt(9)
                cq_run.font.color.rgb = RGBColor(0x71, 0x80, 0x96)
                if ct.get("source_title"):
                    src_run = cq.add_run(f" \u2014 {ct['source_title']}")
                    src_run.font.size = Pt(8.5)
                    src_run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
            # Advisor annotation
            if step.get("advisor_annotation"):
                ap = doc.add_paragraph()
                ap.paragraph_format.left_indent = Inches(0.5)
                ar = ap.add_run(f"Rådgivarens kommentar: {step['advisor_annotation']}")
                ar.font.size = Pt(9)
                ar.font.color.rgb = RGBColor(0x2C, 0x52, 0x82)
            if step.get("conclusion"):
                cp = doc.add_paragraph()
                cp.paragraph_format.left_indent = Inches(0.3)
                cr = cp.add_run(step["conclusion"])
                cr.italic = True

        assumptions = rec.get("assumptions") or []
        if assumptions:
            doc.add_heading("Antaganden", level=3)
            for a in assumptions:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.3)
                bold_run = p.add_run(a["assumption"])
                bold_run.bold = True
                gp = doc.add_paragraph()
                gp.paragraph_format.left_indent = Inches(0.5)
                gp.add_run("Grund: ").bold = True
                gp.add_run(a["basis"])
                ip = doc.add_paragraph()
                ip.paragraph_format.left_indent = Inches(0.5)
                ip.add_run("Konsekvens om fel: ").bold = True
                ip.add_run(a["impact_if_wrong"])

        if evidences:
            doc.add_heading("Bevisunderlag", level=3)
            for ev in evidences:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.3)
                type_run = p.add_run(f"[{ev['source_type'].upper()}] ")
                type_run.font.size = Pt(8.5)
                type_run.font.color.rgb = RGBColor(0x71, 0x80, 0x96)
                ref_run = p.add_run(ev["source_reference"])
                ref_run.bold = True
                snip_p = doc.add_paragraph()
                snip_p.paragraph_format.left_indent = Inches(0.5)
                snip_p.add_run(ev["content_snippet"])
                rel_p = doc.add_paragraph()
                rel_p.paragraph_format.left_indent = Inches(0.5)
                rel_run = rel_p.add_run(ev["relevance_explanation"])
                rel_run.font.size = Pt(9)
                rel_run.font.color.rgb = RGBColor(0x71, 0x80, 0x96)

        # --- 6. Lämplighetsbedömning ---
        add_section(6, "Lämplighetsbedömning")
        score = rec.get("suitability_score")
        if score is not None:
            try:
                score_pct = f"{float(score) * 100:.0f}%"
            except (ValueError, TypeError):
                score_pct = str(score)
            doc.add_paragraph(f"Lämplighetspoäng: {score_pct}")
        assessment = "Denna rekommendation har bedömts mot klientens riskprofil"
        if client.get("risk_profile"):
            assessment += f" ({client['risk_profile']})"
        assessment += ", ekonomiska situation"
        if client.get("annual_income"):
            assessment += f" (årsinkomst {client['annual_income']:,.0f} SEK)"
        assessment += ", och uttalade mål"
        if client.get("desired_retirement_age"):
            assessment += f" (önskad pensionsålder {client['desired_retirement_age']} år)"
        assessment += "."
        doc.add_paragraph(assessment)

        # --- 7. Kostnadsinformation ---
        add_section(7, "Kostnadsinformation")
        cost_text = None
        for step in rec["reasoning_chain"]:
            if "Kostnadsinformation" in step.get("description", ""):
                cost_text = step["conclusion"]
                break
        if cost_text:
            doc.add_paragraph(cost_text)

        if scenarios:
            has_fee = any(s.get("projected_outcome", {}).get("annual_fee") is not None for s in scenarios)
            has_cost = any(s.get("projected_outcome", {}).get("total_cost") is not None for s in scenarios)
            if has_fee or has_cost:
                cols = ["Alternativ"]
                if has_fee:
                    cols.append("Årlig avgift (%)")
                if has_cost:
                    cols.append("Total kostnad (SEK)")
                table = doc.add_table(rows=1, cols=len(cols))
                table.style = "Light List Accent 1"
                for i, c in enumerate(cols):
                    table.rows[0].cells[i].text = c
                for s in scenarios:
                    row = table.add_row()
                    row.cells[0].text = s["name"]
                    ci = 1
                    if has_fee:
                        fee = s.get("projected_outcome", {}).get("annual_fee")
                        try:
                            row.cells[ci].text = f"{float(fee)}%" if fee is not None else "—"
                        except (ValueError, TypeError):
                            row.cells[ci].text = str(fee) if fee is not None else "—"
                        ci += 1
                    if has_cost:
                        cost = s.get("projected_outcome", {}).get("total_cost")
                        try:
                            row.cells[ci].text = f"{float(cost):,.0f} SEK" if cost is not None else "—"
                        except (ValueError, TypeError):
                            row.cells[ci].text = str(cost) if cost is not None else "—"

        # --- 8. Intressekonflikter ---
        add_section(8, "Intressekonflikter")
        conflict_text = None
        for step in rec["reasoning_chain"]:
            if "Intressekonflikt" in step.get("description", ""):
                conflict_text = step["conclusion"]
                break
        doc.add_paragraph(
            conflict_text
            or "Inga intressekonflikter har identifierats i samband med denna rekommendation."
        )

        # --- 9. Rådgivarinformation ---
        add_section(9, "Rådgivarinformation")
        add_kv_table([
            ("Rådgivare", advisor["name"]),
            ("E-post", advisor["email"]),
            ("Rådgivningsdatum", generated_at),
            ("Ärendenummer", case["id"]),
            ("Rekommendationsversion", str(rec["version"])),
        ])

        # --- Footer ---
        footer_para = doc.add_paragraph()
        footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer_run = footer_para.add_run(
            f"\nÄrende {case['id'][:8]}… · Version {rec['version']} · {generated_at} · Genererat av Aetherion"
        )
        footer_run.font.size = Pt(8)
        footer_run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)

        doc.save(str(output_path))
